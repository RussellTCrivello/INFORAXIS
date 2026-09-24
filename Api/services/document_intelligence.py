"""Shared document-intelligence helpers for review and batch extraction.

This module contains data-only operations that are used by search review,
selected-file exports, and the files workspace. It deliberately does not open
paths supplied by a browser: callers resolve file paths from ``paths`` rows
first, and these helpers only read the path attached to that trusted record.
"""

from __future__ import annotations

import re
from collections import defaultdict
from pathlib import Path
from typing import Any, Dict, Iterable, List, Sequence

MAX_FIRST_PAGE_CHARS = 12_000
MAX_CONTACT_CONTEXT_CHARS = 180

_EMAIL_RE = re.compile(
    r"(?<![A-Z0-9.!#$%&'*+/=?^_`{|}~-])"
    r"[A-Z0-9.!#$%&'*+/=?^_`{|}~-]+@"
    r"(?:[A-Z0-9](?:[A-Z0-9-]{0,61}[A-Z0-9])?\.)+[A-Z]{2,63}\b",
    re.IGNORECASE,
)
_URL_RE = re.compile(r"\b(?:https?://|www\.)[^\s<>\"']+", re.IGNORECASE)
_URL_TRAILING_PUNCTUATION = ".,;:!?)]}"
_TOKEN_RE = re.compile(r"[\w]{2,}", re.UNICODE)
_PAGE_LOCATION_RE = re.compile(r"^\s*Page\s+(\d+)\b", re.IGNORECASE | re.MULTILINE)
_SHEET_LOCATION_RE = re.compile(r"^\s*Sheet:\s*([^|\n]+)", re.IGNORECASE | re.MULTILINE)
_SLIDE_LOCATION_RE = re.compile(r"^\s*Slide\s+(\d+)\b", re.IGNORECASE | re.MULTILINE)


def normalize_file_ids(values: Any, *, max_files: int = 200) -> List[int]:
    """Validate a bounded list of positive file IDs, preserving request order."""
    if not isinstance(values, (list, tuple, set)):
        raise ValueError("file_ids must be a list")
    if len(values) > max_files:
        raise ValueError(f"Select no more than {max_files} files at a time")

    ids: List[int] = []
    seen = set()
    for value in values:
        text = str(value).strip()
        if not text.isdigit() or int(text) < 1:
            raise ValueError("file_ids must contain positive whole numbers")
        file_id = int(text)
        if file_id not in seen:
            ids.append(file_id)
            seen.add(file_id)
        if len(ids) > max_files:
            raise ValueError(f"Select no more than {max_files} files at a time")

    if not ids:
        raise ValueError("Select at least one file")
    return ids


def _text_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bytes):
        return value.decode("utf-8", errors="replace")
    return str(value)


def spreadsheet_safe_text(value: Any) -> str:
    """Neutralize formula-leading text before CSV/XLSX export.

    Spreadsheet applications may evaluate cells beginning with ``=``, ``+``,
    ``-`` or ``@`` (also after whitespace/control characters). Prefix those
    values with an apostrophe so an extracted filename or context stays text.
    """
    text = _text_value(value)
    if re.match(r"^[\s\x00-\x1f]*[=+@-]", text):
        return "'" + text
    return text


def _bounded_text(text: str, limit: int = MAX_FIRST_PAGE_CHARS) -> str:
    """Apply a safe output ceiling and prefer a logical page separator."""
    if len(text) <= limit:
        return text.strip()
    first_page = text.split("\f", 1)[0]
    if first_page and len(first_page) < len(text):
        return first_page[:limit].strip()
    return text[:limit].rstrip()


def first_page_text(
    *, file_name: str, file_path: str | None, extracted_text: Any
) -> Dict[str, Any]:
    """Extract the first page/preview segment without changing stored content.

    PDFs use their physical first page. DOCX files stop at an explicit page
    break when present (DOCX pagination otherwise depends on the renderer, so
    their first preview segment is bounded extracted text). XLSX files use the
    first worksheet's first rows. Plain-text and legacy formats use the first
    form-feed-delimited page or a bounded opening segment.
    """
    name = str(file_name or "")
    extension = Path(name).suffix.lower()
    path = Path(file_path) if file_path and "::" not in str(file_path) else None
    if path is not None and not path.is_absolute():
        # The DB path is trusted as an identifier, but relative source paths
        # should not be resolved against the web worker's current directory.
        path = None

    if extension == ".pdf" and path and path.is_file():
        try:
            # PyMuPDF is the project's supported PDF engine. Prefer it so this
            # feature works in the same install profile as PDF ingestion.
            import fitz

            with fitz.open(str(path)) as document:
                if len(document):
                    text = document.load_page(0).get_text("text") or ""
                    if text.strip():
                        return {"text": _bounded_text(text), "method": "pdf-first-page"}
        except Exception:
            pass
        try:
            # Keep a compatibility fallback for installations that provide
            # pypdf/PyPDF2 but not the optional PyMuPDF extra.
            from PyPDF2 import PdfReader

            reader = PdfReader(str(path))
            if reader.pages:
                text = reader.pages[0].extract_text() or ""
                if text.strip():
                    return {"text": _bounded_text(text), "method": "pdf-first-page"}
        except Exception:
            # OCR/text fallback below is still useful for damaged, scanned or
            # encrypted PDFs; never claim a physical-page extract if empty.
            pass

    if extension == ".docx" and path and path.is_file():
        try:
            from docx import Document

            document = Document(str(path))
            parts: List[str] = []
            for paragraph in document.paragraphs:
                parts.append(paragraph.text)
                has_page_break = bool(
                    paragraph._p.xpath('.//w:br[@w:type="page"]')
                    or paragraph._p.xpath(".//w:lastRenderedPageBreak")
                )
                joined = "\n".join(parts)
                if has_page_break or len(joined) >= MAX_FIRST_PAGE_CHARS:
                    break
            text = "\n".join(parts)
            if text.strip():
                return {"text": _bounded_text(text), "method": "docx-first-segment"}
        except Exception:
            pass

    if extension == ".xlsx" and path and path.is_file():
        try:
            from openpyxl import load_workbook

            workbook = load_workbook(str(path), read_only=True, data_only=True)
            try:
                sheet = workbook.worksheets[0] if workbook.worksheets else None
                if sheet is not None:
                    rows = []
                    for row in sheet.iter_rows(max_row=100, values_only=True):
                        values = ["" if cell is None else str(cell) for cell in row]
                        rows.append("\t".join(values).rstrip())
                        if sum(map(len, rows)) >= MAX_FIRST_PAGE_CHARS:
                            break
                    text = "\n".join(rows)
                    if text.strip():
                        return {"text": _bounded_text(text), "method": "xlsx-first-sheet"}
            finally:
                workbook.close()
        except Exception:
            pass

    text = _text_value(extracted_text)
    logical_page = text.split("\f", 1)[0]
    return {
        "text": _bounded_text(logical_page),
        "method": "extracted-text-preview" if text else "unavailable",
    }


def _location_for_offset(source: str, start: int) -> str:
    """Best-effort page, slide, worksheet or line reference for an offset."""
    prefix = source[:start]
    line_number = prefix.count("\n") + 1
    candidates = []
    for pattern, kind in (
        (_PAGE_LOCATION_RE, "page"),
        (_SHEET_LOCATION_RE, "sheet"),
        (_SLIDE_LOCATION_RE, "slide"),
    ):
        for match in pattern.finditer(prefix):
            value = match.group(1).strip()
            candidates.append((match.start(), kind, value))
    if not candidates:
        return f"Line {line_number}"
    _offset, kind, value = max(candidates, key=lambda item: item[0])
    if kind == "page":
        return f"Page {value}, line {line_number}"
    if kind == "slide":
        return f"Slide {value}, line {line_number}"
    return f"Sheet {value}, line {line_number}"


def extract_contact_occurrences(
    text: Any, *, file_id: int, file_name: str
) -> List[Dict[str, Any]]:
    """Return unique email/URL values per file with counts and first context."""
    source = _text_value(text)
    found: Dict[tuple[str, str], Dict[str, Any]] = {}

    matches = []
    matches.extend((match.start(), "email", match.group(0)) for match in _EMAIL_RE.finditer(source))
    for match in _URL_RE.finditer(source):
        value = match.group(0).rstrip(_URL_TRAILING_PUNCTUATION)
        if value:
            matches.append((match.start(), "url", value))

    for start, kind, value in sorted(matches, key=lambda item: item[0]):
        key = (kind, value.casefold())
        item = found.get(key)
        if item is None:
            left = max(0, start - MAX_CONTACT_CONTEXT_CHARS // 2)
            right = min(len(source), start + len(value) + MAX_CONTACT_CONTEXT_CHARS // 2)
            context = re.sub(r"\s+", " ", source[left:right]).strip()
            item = {
                "file_id": file_id,
                "file_name": file_name,
                "kind": kind,
                "value": value,
                "occurrences": 0,
                "first_line": source.count("\n", 0, start) + 1,
                "location": _location_for_offset(source, start),
                "context": context,
            }
            found[key] = item
        item["occurrences"] += 1

    return list(found.values())


def _tokens(text: Any) -> set[str]:
    source = _text_value(text).lower()
    # A fixed text sample makes grouping time bounded even for very large files.
    return set(_TOKEN_RE.findall(source[:100_000]))


def group_similar_documents(
    documents: Sequence[Dict[str, Any]], *, threshold: float = 0.32
) -> List[Dict[str, Any]]:
    """Cluster the supplied page of result documents by token-set similarity.

    This is a deterministic, bounded Jaccard grouping for an explicitly
    requested result page. It is not presented as a forensic duplicate hash or
    a corpus-wide similarity index; exact duplicate suppression uses hashes in
    the search query separately.
    """
    if not 0.05 <= float(threshold) <= 0.95:
        raise ValueError("threshold must be between 0.05 and 0.95")

    items = list(documents)
    if len(items) > 100:
        raise ValueError("Similarity grouping is limited to 100 documents per page")
    token_sets = [_tokens(item.get("content", "")) for item in items]
    parent = list(range(len(items)))
    rank = [0] * len(items)
    edges: Dict[tuple[int, int], float] = {}

    def find(index: int) -> int:
        while parent[index] != index:
            parent[index] = parent[parent[index]]
            index = parent[index]
        return index

    def union(left: int, right: int) -> None:
        root_left, root_right = find(left), find(right)
        if root_left == root_right:
            return
        if rank[root_left] < rank[root_right]:
            root_left, root_right = root_right, root_left
        parent[root_right] = root_left
        if rank[root_left] == rank[root_right]:
            rank[root_left] += 1

    for left in range(len(items)):
        if not token_sets[left]:
            continue
        for right in range(left + 1, len(items)):
            if not token_sets[right]:
                continue
            intersection = len(token_sets[left] & token_sets[right])
            union_size = len(token_sets[left] | token_sets[right])
            score = intersection / union_size if union_size else 0.0
            if score >= threshold:
                edges[(left, right)] = score
                union(left, right)

    clusters: Dict[int, List[int]] = defaultdict(list)
    for index in range(len(items)):
        clusters[find(index)].append(index)

    ordered_clusters = sorted(clusters.values(), key=lambda indexes: indexes[0])
    result: List[Dict[str, Any]] = []
    group_number = 0
    for indexes in ordered_clusters:
        group_number += 1
        ids = [items[index].get("id") for index in indexes]
        similarities = [
            score for (left, right), score in edges.items()
            if left in indexes and right in indexes
        ]
        result.append({
            "group_id": group_number,
            "file_ids": ids,
            "size": len(ids),
            "similarity": round(sum(similarities) / len(similarities), 4)
            if similarities else None,
        })
    return result


__all__ = [
    "MAX_FIRST_PAGE_CHARS",
    "extract_contact_occurrences",
    "first_page_text",
    "group_similar_documents",
    "normalize_file_ids",
    "spreadsheet_safe_text",
]
