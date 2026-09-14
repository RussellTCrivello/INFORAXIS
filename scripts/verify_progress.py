#!/usr/bin/env python3
"""Progress-accounting diagnostic: prints the live timeline and checks the contract.

The regression tests in ``tests/unit/test_nested_progress_accounting.py`` assert
the contract; this script *shows* it. It drives the real engine over a nested
workload and prints the snapshot stream the job system persists, so a "0% then
100%" regression is visible as a flat line rather than as a red test name.

Workloads exercised (see the defect report's verification matrix):
    flat        top-level files only
    archive     one ZIP with several members
    nested      recursively nested ZIPs
    email       an .eml with attachments
    office      a DOCX with an embedded image
    mixed       email -> ZIP -> DOCX -> embedded image
    unsupported supported + unsupported + empty objects
    single      one container through process_single_file
    all         every workload above (default)

Usage:
    python scripts/verify_progress.py                  # all workloads
    python scripts/verify_progress.py --case mixed     # one workload
    python scripts/verify_progress.py --path /some/dir # a real directory
    python scripts/verify_progress.py --delay 0.0      # no artificial delay

Exit code 0 = every contract check passed, 1 = at least one failed.
"""
import argparse
import base64
import io
import os
import shutil
import sys
import tempfile
import threading
import time
import zipfile
from pathlib import Path

PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

BODY = "lorem ipsum dolor sit amet consectetur " * 40

CONTRACT = """\
contract:
  denominator >= real workload          every discovered unit is counted
  files_done == total_files             all of it reached a terminal state
  files_pending == 0                    nothing left queued
  intermediate percentages emitted      the bar moved while work was running
  percent monotonic                     the bar never retreated
  total monotonic                       discovery only ever grew the workload
  100% only at the end                  no premature completion
  nested children attributed            descendants keep their container
"""


# ---------------------------------------------------------------------------
# Observation
# ---------------------------------------------------------------------------
class Recorder:
    """Captures every snapshot pushed through progress_callback, plus samples.

    The callback gives an exact event trace; the sampler approximates what a
    2 s polling frontend would actually have seen between those events.
    """

    def __init__(self, sample_interval=0.05):
        self.events = []
        self.samples = []
        self.t0 = None
        self._stop = threading.Event()
        self._interval = sample_interval
        self._thread = None
        self._reader = None

    def attach(self, reader):
        self._reader = reader
        reader.progress_callback = self._on_snapshot

    def _on_snapshot(self, snapshot):
        self.events.append((round(time.time() - self.t0, 3), dict(snapshot)))

    def start(self):
        self.t0 = time.time()
        self._thread = threading.Thread(target=self._sample, daemon=True)
        self._thread.start()

    def _sample(self):
        while not self._stop.wait(self._interval):
            try:
                self.samples.append(
                    (round(time.time() - self.t0, 3), self._reader.get_live_progress())
                )
            except Exception:
                pass

    def stop(self):
        self._stop.set()
        if self._thread:
            self._thread.join(timeout=2)

    @staticmethod
    def _series(rows, key):
        return [r[1].get(key, 0) for r in rows]

    def percents(self, rows=None):
        return self._series(rows if rows is not None else self.events, "percent")

    def totals(self, rows=None):
        return self._series(rows if rows is not None else self.events, "total_files")


# ---------------------------------------------------------------------------
# Workload builders
# ---------------------------------------------------------------------------
def _png_bytes():
    try:
        from PIL import Image
        buf = io.BytesIO()
        Image.new("RGB", (60, 30), (200, 40, 40)).save(buf, format="PNG")
        return buf.getvalue()
    except Exception:
        import struct
        import zlib

        def chunk(tag, data):
            body = tag + data
            return struct.pack(">I", len(data)) + body + struct.pack(">I", zlib.crc32(body))

        raw = b"".join(b"\x00" + b"\xc8\x28\x28" * 2 for _ in range(2))
        return (b"\x89PNG\r\n\x1a\n"
                + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
                + chunk(b"IDAT", zlib.compress(raw))
                + chunk(b"IEND", b""))


def _docx_bytes():
    try:
        import docx
    except ImportError:
        return None
    document = docx.Document()
    document.add_paragraph("Body text that must survive the container merge.")
    document.add_picture(io.BytesIO(_png_bytes()))
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def _zip_bytes(entries):
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        for name, payload in entries.items():
            zf.writestr(name, payload)
    return buf.getvalue()


def _eml_bytes(attachments):
    boundary = "PROGRESSBOUNDARY"
    buf = io.StringIO()
    buf.write("From: sender@example.com\r\nTo: recipient@example.com\r\n"
              "Subject: progress accounting\r\nMIME-Version: 1.0\r\n"
              f'Content-Type: multipart/mixed; boundary="{boundary}"\r\n\r\n')
    buf.write(f"--{boundary}\r\nContent-Type: text/plain; charset=utf-8\r\n\r\n"
              "Message body that must survive alongside the attachments.\r\n")
    for filename, ctype, payload in attachments:
        buf.write(f"--{boundary}\r\nContent-Type: {ctype}\r\n"
                  "Content-Transfer-Encoding: base64\r\n"
                  f'Content-Disposition: attachment; filename="{filename}"\r\n\r\n')
        buf.write(base64.encodebytes(payload).decode() + "\r\n")
    buf.write(f"--{boundary}--\r\n")
    return buf.getvalue().encode()


def build_flat(root):
    for i in range(4):
        (root / f"doc{i}.txt").write_text(BODY)
    return 4


def build_archive(root):
    (root / "bundle.zip").write_bytes(
        _zip_bytes({f"m{i}.txt": BODY.encode() for i in range(6)})
    )
    (root / "top.txt").write_text(BODY)
    return 8  # top.txt + bundle.zip + 6 members


def build_nested(root):
    level3 = _zip_bytes({f"deep{i}.txt": BODY.encode() for i in range(2)})
    level2 = _zip_bytes({
        **{f"mid{i}.txt": BODY.encode() for i in range(2)},
        "level3.zip": level3,
    })
    (root / "level1.zip").write_bytes(_zip_bytes({"level2.zip": level2}))
    (root / "top.txt").write_text(BODY)
    return 8  # top + l1 + l2 + l3 + 2 mid + 2 deep


def build_email(root):
    (root / "message.eml").write_bytes(_eml_bytes([
        (f"att{i}.txt", "text/plain", BODY.encode()) for i in range(3)
    ]))
    return 4  # eml + 3 attachments


def build_office(root):
    payload = _docx_bytes()
    if payload is None:
        print("  (python-docx not installed; skipping office case)")
        return None
    (root / "report.docx").write_bytes(payload)
    (root / "top.txt").write_text(BODY)
    return 3  # top + docx + embedded image (at least)


def build_mixed(root):
    payload = _docx_bytes()
    if payload is None:
        print("  (python-docx not installed; skipping mixed case)")
        return None
    archive = _zip_bytes({
        "inner.docx": payload,
        **{f"side{i}.txt": BODY.encode() for i in range(2)},
    })
    (root / "thread.eml").write_bytes(
        _eml_bytes([("payload.zip", "application/zip", archive)])
    )
    return 6  # eml + zip + docx + 2 sides + >=1 embedded image


def build_unsupported(root):
    (root / "good.txt").write_text(BODY)
    (root / "mystery.zzzq").write_bytes(b"\x00\x01\x02not a format\x03" * 60)
    (root / "empty.bin").write_bytes(b"")
    return 3


def build_single(root):
    (root / "bundle.zip").write_bytes(
        _zip_bytes({f"m{i}.txt": BODY.encode() for i in range(5)})
    )
    return 6  # zip + 5 members


CASES = {
    "flat": (build_flat, "folder"),
    "archive": (build_archive, "folder"),
    "nested": (build_nested, "folder"),
    "email": (build_email, "folder"),
    "office": (build_office, "folder"),
    "mixed": (build_mixed, "folder"),
    "unsupported": (build_unsupported, "folder"),
    "single": (build_single, "file"),
}


# ---------------------------------------------------------------------------
# Contract checks
# ---------------------------------------------------------------------------
class Report:
    def __init__(self):
        self.passed = []
        self.failed = []

    def check(self, name, condition, detail=""):
        (self.passed if condition else self.failed).append((name, detail))
        mark = "PASS" if condition else "FAIL"
        suffix = f"  <- {detail}" if (detail and not condition) else ""
        print(f"  [{mark}] {name}{suffix}")


def verify(final, recorder, expected_total, report, mode):
    print("\n  -- progress timeline (as a polling frontend would see it) --")
    print(f"  {'t(s)':>7} {'percent':>8} {'total':>6} {'done':>5} {'nested':>7} {'running':>8}")
    last = None
    for t, snap in recorder.samples:
        row = (snap.get("percent"), snap.get("total_files"), snap.get("files_done"),
               snap.get("files_nested"), snap.get("in_progress"))
        if row != last:
            print(f"  {t:>7} {row[0]:>8} {row[1]:>6} {row[2]:>5} {row[3]:>7} {row[4]:>8}")
            last = row

    percents = recorder.percents()
    totals = recorder.totals()
    intermediates = [p for p in percents if 0 < p < 100]

    print("\n  -- final snapshot --")
    print(f"  total={final['total_files']} initial={final['files_initial']} "
          f"nested={final['files_nested']} done={final['files_done']} "
          f"percent={final['percent']} containers={final['containers_opened']}")
    print(f"  completed={final['files_completed']} failed={final['files_failed']} "
          f"skipped={final['files_skipped']} unsupported={final['files_unsupported']} "
          f"retryable={final['files_retryable']} pending={final['files_pending']} "
          f"running={final['in_progress']}")
    if final.get("children_by_parent"):
        print("  attribution: " + ", ".join(
            f"{os.path.basename(k)}->{v}" for k, v in final["children_by_parent"].items()
        ))

    print("\n  -- contract --")
    report.check("denominator reaches the real workload",
                 final["total_files"] >= expected_total,
                 f"total={final['total_files']} expected>={expected_total}")
    report.check("every discovered unit reached a terminal state",
                 final["files_done"] == final["total_files"],
                 f"done={final['files_done']} total={final['total_files']}")
    report.check("nothing left pending", final["files_pending"] == 0,
                 f"pending={final['files_pending']}")
    report.check("final percent is 100", final["percent"] == 100,
                 f"percent={final['percent']}")
    report.check("ledger reports complete", final["complete"] is True)
    report.check("progress pushed more than once (not only at the end)",
                 len(recorder.events) > 2, f"{len(recorder.events)} snapshots")
    report.check("intermediate percentages were emitted during the run",
                 bool(intermediates), f"distinct percents={sorted(set(percents))}")
    report.check("percent never retreated",
                 all(b >= a for a, b in zip(percents, percents[1:], strict=False)),
                 f"percents={percents}")
    report.check("total never shrank",
                 all(b >= a for a, b in zip(totals, totals[1:], strict=False)),
                 f"totals={totals}")

    first_full = next((i for i, p in enumerate(percents) if p == 100), None)
    report.check("100% was not reported prematurely",
                 first_full is None or all(p == 100 for p in percents[first_full:]),
                 f"100% first at index {first_full} of {len(percents)}")

    if final["files_nested"] > 0:
        report.check("total grew dynamically as containers opened",
                     len(set(totals)) > 1, f"distinct totals={sorted(set(totals))}")
        report.check("initial vs discovered totals distinguished",
                     final["files_initial"] < final["files_discovered"],
                     f"initial={final['files_initial']} discovered={final['files_discovered']}")
        report.check("nested children attributed to their container",
                     bool(final.get("children_by_parent")),
                     f"attribution={final.get('children_by_parent')}")
    else:
        report.check("flat workload reports a static total",
                     len(set(totals)) == 1, f"distinct totals={sorted(set(totals))}")

    if mode == "file":
        report.check("single-file path publishes a denominator immediately",
                     bool(recorder.events) and recorder.events[0][1]["total_files"] >= 1,
                     f"first snapshot={recorder.events[0][1] if recorder.events else None}")


# ---------------------------------------------------------------------------
def run_case(name, delay, report, target_path=None):
    from pipeline.integrated_reader import IntegratedFileReader

    builder, mode = CASES[name]
    print(f"\n{'=' * 78}\nCASE: {name}  ({'process_single_file' if mode == 'file' else 'process_folder'})\n{'=' * 78}")

    root = Path(tempfile.mkdtemp(prefix=f"prog_{name}_"))
    try:
        if target_path:
            expected_total, mode, entry = None, "folder", Path(target_path)
        else:
            expected_total = builder(root)
            if expected_total is None:
                return
            entry = root / "bundle.zip" if mode == "file" else root

        # Artificial delay so the timeline has resolution. Patched at the engine
        # boundary, never inside the ledger, so accounting is unaffected.
        import pipeline.integrated_reader as ir
        original = ir.main_specify_method_of_reading_the_file

        def delayed(file_info, **kwargs):
            if delay:
                time.sleep(delay)
            return original(file_info, **kwargs)

        ir.main_specify_method_of_reading_the_file = delayed
        try:
            reader = IntegratedFileReader(
                max_workers=2, enable_monitoring=False, enable_storage=False
            )
            recorder = Recorder()
            recorder.attach(reader)
            recorder.start()
            if mode == "file":
                reader.process_single_file(str(entry))
            else:
                reader.process_folder(str(entry))
            recorder.stop()
        finally:
            ir.main_specify_method_of_reading_the_file = original

        final = reader.get_live_progress()
        verify(final, recorder, expected_total or final["total_files"], report, mode)
    finally:
        shutil.rmtree(root, ignore_errors=True)


def main():
    parser = argparse.ArgumentParser(description=__doc__,
                                     formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("--case", choices=sorted(CASES), help="run a single workload")
    parser.add_argument("--path", help="run against a real file or directory instead")
    parser.add_argument("--delay", type=float, default=0.15,
                        help="artificial per-file delay in seconds (default 0.15)")
    args = parser.parse_args()

    print(CONTRACT)
    report = Report()

    if args.path:
        target = Path(args.path)
        mode = "file" if target.is_file() else "folder"
        print(f"\n{'=' * 78}\nCASE: --path {target}  ({mode})\n{'=' * 78}")
        from pipeline.integrated_reader import IntegratedFileReader

        reader = IntegratedFileReader(max_workers=2, enable_monitoring=False,
                                      enable_storage=False)
        recorder = Recorder()
        recorder.attach(reader)
        recorder.start()
        if mode == "file":
            reader.process_single_file(str(target))
        else:
            reader.process_folder(str(target))
        recorder.stop()
        verify(reader.get_live_progress(), recorder, 1, report, mode)
    else:
        names = [args.case] if args.case else list(CASES)
        for name in names:
            run_case(name, args.delay, report)

    print(f"\n{'#' * 78}")
    print(f"# RESULT: {len(report.passed)} passed, {len(report.failed)} failed")
    for name, detail in report.failed:
        print(f"#   FAILED: {name}  ({detail})")
    print("#" * 78)
    return 1 if report.failed else 0


if __name__ == "__main__":
    sys.exit(main())
